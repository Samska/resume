#!/usr/bin/env python3
"""Build ATS-friendly Markdown, HTML, PDF, and DOCX resumes from JSON."""

from __future__ import annotations

import argparse
import html
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageTemplate,
    Paragraph,
    Spacer,
)


ROOT = Path(__file__).resolve().parents[1]
CONTENT = ROOT / "content"
DIST = ROOT / "dist"
STYLE = ROOT / "styles" / "resume.css"


def load(locale: str) -> dict:
    path = CONTENT / f"resume.{locale}.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    required = {"locale", "name", "contact", "section_labels", "summary", "skills", "experience", "education"}
    missing = required - data.keys()
    if missing:
        raise ValueError(f"{path.name}: missing fields: {', '.join(sorted(missing))}")
    if data["locale"] != locale:
        raise ValueError(f"{path.name}: locale must be {locale}")
    return data


def markdown(data: dict) -> str:
    labels = data["section_labels"]
    out = [f"# {data['name']}", ""]
    out.append(" | ".join(f"[{item['text']}]({item['url']})" for item in data["contact"]))
    out.extend(["", f"## {labels['summary']}", ""])
    out.extend(data["summary"])
    out.extend(["", f"## {labels['skills']}", ""])
    for skill in data["skills"]:
        out.append(f"**{skill['category']}:** {', '.join(skill['items'])}")
    out.extend(["", f"## {labels['experience']}", ""])
    for job in data["experience"]:
        out.append(f"### {job['role']} | {job['company']}")
        out.append(f"{job['dates']} | {job['location']}")
        out.extend(f"- {bullet}" for bullet in job["bullets"])
        out.append("")
    out.extend([f"## {labels['education']}", ""])
    for item in data["education"]:
        out.append(f"### {item['program']} | {item['institution']}")
        out.append(item["dates"])
        out.append("")
    return "\n".join(out).rstrip() + "\n"


def html_document(data: dict, md: str) -> str:
    labels = data["section_labels"]
    css = STYLE.read_text(encoding="utf-8")
    contact = " | ".join(
        f'<a href="{html.escape(item["url"])}">{html.escape(item["text"])}</a>'
        for item in data["contact"]
    )
    blocks = [
        "<!doctype html>",
        f'<html lang="{html.escape(data["locale"])}"><head><meta charset="utf-8">',
        f"<title>{html.escape(data['name'])} Resume</title><style>{css}</style></head><body>",
        f"<h1>{html.escape(data['name'])}</h1>",
        f'<p class="contact">{contact}</p>',
        f"<h2>{html.escape(labels['summary'])}</h2>",
    ]
    blocks.extend(f"<p>{html.escape(text)}</p>" for text in data["summary"])
    blocks.append(f"<h2>{html.escape(labels['skills'])}</h2>")
    blocks.extend(
        f"<p><strong>{html.escape(skill['category'])}:</strong> {html.escape(', '.join(skill['items']))}</p>"
        for skill in data["skills"]
    )
    blocks.append(f"<h2>{html.escape(labels['experience'])}</h2>")
    for job in data["experience"]:
        blocks.append(f"<h3>{html.escape(job['role'])} | {html.escape(job['company'])}</h3>")
        blocks.append(f"<p>{html.escape(job['dates'])} | {html.escape(job['location'])}</p><ul>")
        blocks.extend(f"<li>{html.escape(bullet)}</li>" for bullet in job["bullets"])
        blocks.append("</ul>")
    blocks.append(f"<h2>{html.escape(labels['education'])}</h2>")
    for item in data["education"]:
        blocks.append(f"<h3>{html.escape(item['program'])} | {html.escape(item['institution'])}</h3>")
        blocks.append(f"<p>{html.escape(item['dates'])}</p>")
    blocks.append("</body></html>\n")
    return "\n".join(blocks)


def pdf(data: dict, output: Path) -> None:
    font_dir = Path("/usr/share/fonts/truetype/dejavu")
    regular_font = font_dir / "DejaVuSans.ttf"
    bold_font = font_dir / "DejaVuSans-Bold.ttf"
    if not regular_font.exists() or not bold_font.exists():
        raise FileNotFoundError("DejaVu Sans fonts are required to build the PDF")
    if "ResumeSans" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("ResumeSans", regular_font))
        pdfmetrics.registerFont(TTFont("ResumeSans-Bold", bold_font))

    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "ResumeBody", parent=styles["BodyText"], fontName="ResumeSans", fontSize=9.2,
        leading=11.1, textColor=colors.HexColor("#111111"), spaceAfter=3.5,
    )
    title = ParagraphStyle(
        "ResumeTitle", parent=body, fontName="ResumeSans-Bold", fontSize=20,
        leading=23, alignment=TA_CENTER, spaceAfter=3,
    )
    contact = ParagraphStyle(
        "ResumeContact", parent=body, fontSize=8.7, leading=10.5,
        alignment=TA_CENTER, spaceAfter=7,
    )
    heading = ParagraphStyle(
        "ResumeHeading", parent=body, fontName="ResumeSans-Bold", fontSize=10.6,
        leading=12.2, spaceBefore=6, spaceAfter=3, textTransform="uppercase",
    )
    job = ParagraphStyle(
        "ResumeJob", parent=body, fontName="ResumeSans-Bold", fontSize=9.5,
        leading=11.4, spaceBefore=3.5, spaceAfter=1,
    )
    meta = ParagraphStyle(
        "ResumeMeta", parent=body, fontName="ResumeSans", fontSize=8.5,
        leading=10, spaceAfter=2,
    )
    bullet_style = ParagraphStyle(
        "ResumeBullet", parent=body, fontSize=8.9, leading=10.8, spaceAfter=1.1,
        leftIndent=12, firstLineIndent=-9,
    )

    doc = BaseDocTemplate(
        str(output), pagesize=LETTER, leftMargin=0.58 * inch, rightMargin=0.58 * inch,
        topMargin=0.42 * inch, bottomMargin=0.42 * inch,
        title=f"{data['name']} Resume", author=data["name"],
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates(PageTemplate(id="resume", frames=[frame]))
    story = [Paragraph(html.escape(data["name"]), title)]
    links = " &nbsp;|&nbsp; ".join(
        f'<link href="{html.escape(item["url"])}">{html.escape(item["text"])}</link>'
        for item in data["contact"]
    )
    story.append(Paragraph(links, contact))
    labels = data["section_labels"]
    story.append(Paragraph(html.escape(labels["summary"]), heading))
    story.extend(Paragraph(html.escape(text), body) for text in data["summary"])
    story.append(Paragraph(html.escape(labels["skills"]), heading))
    for skill in data["skills"]:
        story.append(Paragraph(f"<b>{html.escape(skill['category'])}:</b> {html.escape(', '.join(skill['items']))}", body))
    story.append(Paragraph(html.escape(labels["experience"]), heading))
    for item in data["experience"]:
        header = Paragraph(f"{html.escape(item['role'])} | {html.escape(item['company'])}", job)
        metadata = Paragraph(f"{html.escape(item['dates'])} | {html.escape(item['location'])}", meta)
        story.append(KeepTogether([header, metadata]))
        for text in item["bullets"]:
            story.append(Paragraph(f"&#8226;&nbsp; {html.escape(text)}", bullet_style))
        story.append(Spacer(1, 2))
    story.append(Paragraph(html.escape(labels["education"]), heading))
    for item in data["education"]:
        story.append(Paragraph(f"{html.escape(item['program'])} | {html.escape(item['institution'])}", job))
        story.append(Paragraph(html.escape(item["dates"]), meta))
    doc.build(story)


def set_cellless_border(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "B7B7B7")
    borders.append(bottom)


def hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relation_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    props.append(color)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    link.append(run)
    paragraph._p.append(link)


def set_font(run, name: str = "Arial", size: float = 10.0, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def docx(data: dict, output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    document.core_properties.title = f"{data['name']} Resume"
    document.core_properties.author = data["name"]

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.02

    title_style = document.styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(20)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(2)

    h1 = document.styles["Heading 1"]
    h1.font.name = "Arial"
    h1._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    h1.font.size = Pt(10.6)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.space_before = Pt(5)
    h1.paragraph_format.space_after = Pt(2)
    h1.paragraph_format.keep_with_next = True

    role_style = document.styles.add_style("Resume Role", WD_STYLE_TYPE.PARAGRAPH)
    role_style.base_style = normal
    role_style.font.name = "Arial"
    role_style.font.size = Pt(9.5)
    role_style.font.bold = True
    role_style.font.color.rgb = RGBColor(0, 0, 0)
    role_style.paragraph_format.space_before = Pt(3)
    role_style.paragraph_format.space_after = Pt(0)
    role_style.paragraph_format.keep_with_next = True

    name_p = document.add_paragraph(style="Title")
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(name_p.add_run(data["name"]), size=20, bold=True)
    contact_p = document.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.paragraph_format.space_after = Pt(5)
    for index, item in enumerate(data["contact"]):
        if index:
            set_font(contact_p.add_run(" | "), size=8.6)
        hyperlink(contact_p, item["text"], item["url"])

    labels = data["section_labels"]
    for label, payload in (
        (labels["summary"], "summary"),
        (labels["skills"], "skills"),
    ):
        p = document.add_paragraph(label.upper(), style="Heading 1")
        set_cellless_border(p)
        if payload == "summary":
            for text in data["summary"]:
                document.add_paragraph(text)
        else:
            for skill in data["skills"]:
                p = document.add_paragraph()
                set_font(p.add_run(f"{skill['category']}: "), size=9.2, bold=True)
                set_font(p.add_run(", ".join(skill["items"])), size=9.2)

    p = document.add_paragraph(labels["experience"].upper(), style="Heading 1")
    set_cellless_border(p)
    for item in data["experience"]:
        p = document.add_paragraph(style="Resume Role")
        set_font(p.add_run(f"{item['role']} | {item['company']}"), size=9.5, bold=True)
        meta_p = document.add_paragraph()
        meta_p.paragraph_format.space_after = Pt(1)
        meta_p.paragraph_format.keep_with_next = True
        set_font(meta_p.add_run(f"{item['dates']} | {item['location']}"), size=8.5, italic=True)
        for bullet in item["bullets"]:
            bp = document.add_paragraph(style="List Bullet")
            bp.paragraph_format.left_indent = Inches(0.18)
            bp.paragraph_format.first_line_indent = Inches(-0.12)
            bp.paragraph_format.space_after = Pt(0.7)
            bp.paragraph_format.line_spacing = 1.0
            set_font(bp.add_run(bullet), size=8.9)

    p = document.add_paragraph(labels["education"].upper(), style="Heading 1")
    set_cellless_border(p)
    for item in data["education"]:
        p = document.add_paragraph(style="Resume Role")
        set_font(p.add_run(f"{item['program']} | {item['institution']}"), size=9.4, bold=True)
        meta_p = document.add_paragraph()
        set_font(meta_p.add_run(item["dates"]), size=8.5, italic=True)

    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--locale", choices=["pt-BR", "en-US", "all"], default="all")
    args = parser.parse_args()
    locales = ["pt-BR", "en-US"] if args.locale == "all" else [args.locale]
    DIST.mkdir(parents=True, exist_ok=True)
    for locale in locales:
        data = load(locale)
        stem = f"Samuel-Andrade-Resume-{locale}"
        md = markdown(data)
        (DIST / f"{stem}.md").write_text(md, encoding="utf-8")
        (DIST / f"{stem}.html").write_text(html_document(data, md), encoding="utf-8")
        pdf(data, DIST / f"{stem}.pdf")
        docx(data, DIST / f"{stem}.docx")
        print(f"built {stem}")


if __name__ == "__main__":
    main()
